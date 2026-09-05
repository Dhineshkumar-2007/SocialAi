"""Seed all 20 Tamil Nadu universities from Excel into SQLite DB."""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from database.db import get_db, init_db

EXCEL_PATH = "Tamil_Nadu_20_Universities_SocialAI_Onboarding.xlsx"

UNIVERSITY_MAP = {
    # (short_name_or_name_fragment) -> (name, city, description, capacity, email, verified)
}

# Read all rows from Excel to build insertion data
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb["Universities"]

headers = [cell.value for cell in ws[1]]
print("Headers:", headers)

rows = []
for row in ws.iter_rows(min_row=2, values_only=True):
    rows.append(row)

# Print first 3 rows to confirm structure
for r in rows[:3]:
    print(r[:6])

# Initialize DB (creates schema + calls existing seeds if needed)
init_db()

with get_db() as db:
    # Check current count
    count = db.execute("SELECT COUNT(*) FROM universities").fetchone()[0]
    print(f"Current universities count: {count}")

    # Build insert from Excel data
    # Map column names by index based on header
    col_index = {name: i for i, name in enumerate(headers) if name}
    print("Column indices:", col_index)

    # For all 20 rows, insert/update with verified=1
    for r in rows:
        if not r or not r[col_index.get("University_Name", 1)]:
            continue
        name = r[col_index.get("University_Name", 1)]
        short = r[col_index.get("Short_Name", 2)]
        city_col = "City / Main TN Campus"
        city = r[col_index.get(city_col)] if city_col in col_index else None
        institution_type = r[col_index.get("Institution_Type", 3)] if "Institution_Type" in col_index else None
        domain = r[col_index.get("Primary_Domains", 5)] if "Primary_Domains" in col_index else None
        tags = r[col_index.get("Expertise_Tags", 6)] if "Expertise_Tags" in col_index else None
        dept = r[col_index.get("Key_Departments", 7)] if "Key_Departments" in col_index else None
        labs = r[col_index.get("Research_Centres/Labs", 8)] if "Research_Centres/Labs" in col_index else None
        innovation = r[col_index.get("Innovation/Incubation", 9)] if "Innovation/Incubation" in col_index else None
        problem_match = r[col_index.get("Best_Societal_Problem_Matches", 10)] if "Best_Societal_Problem_Matches" in col_index else None
        email = r[col_index.get("Public_Contact_Email", 11)] if "Public_Contact_Email" in col_index else None
        phone = r[col_index.get("Public_Contact_Phone", 12)] if "Public_Contact_Phone" in col_index else None
        website = r[col_index.get("Official_Website", 13)] if "Official_Website" in col_index else None

        description_parts = []
        if institution_type:
            description_parts.append(str(institution_type))
        if city:
            description_parts.append(f"Located in {city}")
        if tags:
            description_parts.append(f"Expertise: {tags}")
        if dept:
            description_parts.append(f"Key departments include {dept}")
        if labs:
            description_parts.append(f"Research labs: {labs}")
        if innovation:
            description_parts.append(f"Innovation/incubation: {innovation}")
        if problem_match:
            description_parts.append(f"Societal problem matches: {problem_match}")
        description = "; ".join(description_parts) if description_parts else "Tamil Nadu higher education institution."

        # Check if exists by name
        existing = db.execute("SELECT id FROM universities WHERE name=?", (name,)).fetchone()
        if existing:
            db.execute(
                "UPDATE universities SET city=?, description=?, verified=1 WHERE id=?",
                (city, description, existing["id"])
            )
            uid = existing["id"]
            print(f"Updated: {name} (id={uid})")
        else:
            cur = db.execute(
                "INSERT INTO universities (name, city, description, capacity, verified, contact_email) VALUES (?, ?, ?, ?, ?, ?)",
                (name, city, description, 10, 1, email)
            )
            uid = cur.lastrowid
            print(f"Inserted: {name} (id={uid}, city={city})")

print("Done seeding 20 universities from Excel.")
